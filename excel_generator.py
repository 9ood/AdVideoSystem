import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

class ExcelGenerator:
    def __init__(self, output_dir):
        self.output_dir = output_dir
    
    def generate(self, global_info, scenes, condensed_script=None, two_part_script=None, creative_core=None, clean_mode=False):
        wb = Workbook()
        
        ws_global = wb.active
        ws_global.title = "全局信息"
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        ws_global['A1'] = '项目'
        ws_global['B1'] = '内容'
        ws_global['A1'].fill = header_fill
        ws_global['B1'].fill = header_fill
        ws_global['A1'].font = header_font
        ws_global['B1'].font = header_font
        
        global_data = [
            ['视频总时长', f"{scenes[-1]['end_time']:.2f}秒" if scenes else '0秒'],
            ['总镜头数', len(scenes)],
            ['整体风格', global_info.get('style', '')],
            ['主角描述', global_info.get('main_character', '')],
            ['场景描述', global_info.get('main_scene', '')],
            ['首帧图片', './frames/global_first.jpg'],
            ['尾帧图片', './frames/global_last.jpg']
        ]
        
        for i, row_data in enumerate(global_data, start=2):
            ws_global[f'A{i}'] = row_data[0]
            ws_global[f'B{i}'] = row_data[1]
        
        ws_global.column_dimensions['A'].width = 20
        ws_global.column_dimensions['B'].width = 60

        if creative_core:
            ws_creative = wb.create_sheet(title="创意内核")

            ws_creative['A1'] = '项目'
            ws_creative['B1'] = '内容'
            ws_creative['A1'].fill = header_fill
            ws_creative['B1'].fill = header_fill
            ws_creative['A1'].font = header_font
            ws_creative['B1'].font = header_font

            creative_core_data = [
                ['一句话创意内核', creative_core.get('one_sentence', '')],
                ['广告类型', creative_core.get('ad_type', '')],
                ['开头钩子', creative_core.get('opening_hook', '')],
                ['故事线', creative_core.get('storyline', '')],
                ['核心冲突/看点', creative_core.get('conflict', '')],
                ['App在故事里的角色', creative_core.get('product_role', '')],
                ['App界面出现方式', creative_core.get('app_ui', '')],
                ['证明App有用的方式', creative_core.get('proof_method', '')],
                ['关键镜头套路', creative_core.get('visual_structure', '')],
                ['情绪变化', creative_core.get('emotion_curve', '')],
                ['翻拍必须保留的骨架', creative_core.get('must_preserve', '')],
                ['翻拍可以替换的元素', creative_core.get('replaceable', '')],
                ['最容易被错误套模板的地方', creative_core.get('template_risk', '')],
                ['西西魔法钢琴翻拍建议', creative_core.get('remake_guidance', '')],
            ]

            for i, row_data in enumerate(creative_core_data, start=2):
                ws_creative[f'A{i}'] = row_data[0]
                ws_creative[f'B{i}'] = row_data[1]
                ws_creative[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top')
                ws_creative.row_dimensions[i].height = 45

            ws_creative.column_dimensions['A'].width = 28
            ws_creative.column_dimensions['B'].width = 100
        
        if condensed_script:
            ws_condensed = wb.create_sheet(title="浓缩脚本")
            
            ws_condensed['A1'] = '项目'
            ws_condensed['B1'] = '内容'
            ws_condensed['A1'].fill = header_fill
            ws_condensed['B1'].fill = header_fill
            ws_condensed['A1'].font = header_font
            ws_condensed['B1'].font = header_font
            
            if not clean_mode:
                ws_condensed['C1'] = '产品图片'
                ws_condensed['C1'].fill = header_fill
                ws_condensed['C1'].font = header_font
            
            condensed_data = [
                ['视频时长', condensed_script.get('duration', '10-15秒')],
                ['核心故事', condensed_script.get('core_story', '')],
                ['创意继承检查', condensed_script.get('creative_inheritance_check', '')],
                ['关键画面1 (0-5秒)', condensed_script.get('key_scene_1', '')],
                ['关键画面2 (5-10秒)', condensed_script.get('key_scene_2', '')],
                ['关键画面3 (10-15秒)', condensed_script.get('key_scene_3', '')],
                ['完整Sora提示词', condensed_script.get('complete_prompt', '')],
                ['中文提示词(Seedance 2.0)', condensed_script.get('seedance_prompt', '')]
            ]
            
            for i, row_data in enumerate(condensed_data, start=2):
                ws_condensed[f'A{i}'] = row_data[0]
                ws_condensed[f'B{i}'] = row_data[1]
                ws_condensed[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top')
            
            if not clean_mode:
                product_images = condensed_script.get('product_images', [])
                if product_images:
                    row_idx = 9
                    ws_condensed.row_dimensions[row_idx].height = max(100, len(product_images) * 80)
                    
                    for img_idx, img_path in enumerate(product_images):
                        if os.path.exists(img_path):
                            try:
                                img = OpenpyxlImage(img_path)
                                img.width = 150
                                img.height = 150
                                
                                cell_position = f'C{row_idx}'
                                img.anchor = cell_position
                                ws_condensed.add_image(img)
                                
                                ws_condensed[cell_position] = os.path.basename(img_path)
                            except Exception as e:
                                print(f"无法插入图片 {img_path}: {e}")
            
            ws_condensed.column_dimensions['A'].width = 25
            ws_condensed.column_dimensions['B'].width = 100
            if not clean_mode:
                ws_condensed.column_dimensions['C'].width = 30
            
            ws_condensed.row_dimensions[7].height = 80
            ws_condensed.row_dimensions[8].height = 100
        
        if two_part_script:
            ws_two_part = wb.create_sheet(title="两段式脚本")
            
            ws_two_part['A1'] = '项目'
            ws_two_part['B1'] = '内容'
            ws_two_part['A1'].fill = header_fill
            ws_two_part['B1'].fill = header_fill
            ws_two_part['A1'].font = header_font
            ws_two_part['B1'].font = header_font
            
            if not clean_mode:
                ws_two_part['C1'] = '首尾帧'
                ws_two_part['D1'] = '产品图片'
                ws_two_part['C1'].fill = header_fill
                ws_two_part['D1'].fill = header_fill
                ws_two_part['C1'].font = header_font
                ws_two_part['D1'].font = header_font
            
            two_part_data = [
                ['核心故事', two_part_script.get('core_story', '')],
                ['上集提示词', two_part_script.get('part1_seedance_prompt', '')],
                ['下集提示词', two_part_script.get('part2_seedance_prompt', '')]
            ]
            
            for i, row_data in enumerate(two_part_data, start=2):
                ws_two_part[f'A{i}'] = row_data[0]
                ws_two_part[f'B{i}'] = row_data[1]
                ws_two_part[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top')
            
            if not clean_mode:
                part1_images = two_part_script.get('part1_product_images', [])
                if part1_images:
                    row_idx = 3
                    ws_two_part.row_dimensions[row_idx].height = max(100, len(part1_images) * 80)
                    
                    for img_idx, img_path in enumerate(part1_images):
                        if os.path.exists(img_path):
                            try:
                                img = OpenpyxlImage(img_path)
                                img.width = 150
                                img.height = 150
                                
                                cell_position = f'D{row_idx}'
                                img.anchor = cell_position
                                ws_two_part.add_image(img)
                                
                                ws_two_part[cell_position] = os.path.basename(img_path)
                            except Exception as e:
                                print(f"无法插入图片 {img_path}: {e}")
                
                part2_images = two_part_script.get('part2_product_images', [])
                if part2_images:
                    row_idx = 4
                    ws_two_part.row_dimensions[row_idx].height = max(100, len(part2_images) * 80)
                    
                    for img_idx, img_path in enumerate(part2_images):
                        if os.path.exists(img_path):
                            try:
                                img = OpenpyxlImage(img_path)
                                img.width = 150
                                img.height = 150
                                
                                cell_position = f'D{row_idx}'
                                img.anchor = cell_position
                                ws_two_part.add_image(img)
                                
                                ws_two_part[cell_position] = os.path.basename(img_path)
                            except Exception as e:
                                print(f"无法插入图片 {img_path}: {e}")
            
            ws_two_part.column_dimensions['A'].width = 30
            ws_two_part.column_dimensions['B'].width = 100
            if not clean_mode:
                ws_two_part.column_dimensions['C'].width = 20
                ws_two_part.column_dimensions['D'].width = 30
            
            ws_two_part.row_dimensions[3].height = 100
            ws_two_part.row_dimensions[4].height = 100
        
        ws_scenes = wb.create_sheet(title="镜头详情")
        
        if clean_mode:
            headers = ['镜头编号', '开始时间(秒)', '结束时间(秒)', '时长(秒)', '镜头类型', 'Sora提示词(场景描述)', '台词脚本', '完整Sora提示词(可直接使用)', '中文提示词(Seedance 2.0)']
        else:
            headers = ['镜头编号', '开始时间(秒)', '结束时间(秒)', '时长(秒)', '镜头类型', 'Sora提示词(场景描述)', '台词脚本', '完整Sora提示词(可直接使用)', '中文提示词(Seedance 2.0)', '产品图片', '关键帧路径']
        
        for col, header in enumerate(headers, start=1):
            cell = ws_scenes.cell(row=1, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for i, scene in enumerate(scenes, start=2):
            ws_scenes.cell(row=i, column=1, value=scene['scene_number'])
            ws_scenes.cell(row=i, column=2, value=f"{scene['start_time']:.2f}")
            ws_scenes.cell(row=i, column=3, value=f"{scene['end_time']:.2f}")
            ws_scenes.cell(row=i, column=4, value=f"{scene['duration']:.2f}")
            
            shot_type = self.extract_shot_type(scene['prompt'])
            ws_scenes.cell(row=i, column=5, value=shot_type)
            
            ws_scenes.cell(row=i, column=6, value=scene['prompt'])
            
            script = scene.get('script', '')
            ws_scenes.cell(row=i, column=7, value=script)
            
            complete_prompt = scene.get('complete_prompt', '')
            ws_scenes.cell(row=i, column=8, value=complete_prompt)
            
            seedance_prompt_data = scene.get('seedance_prompt', '')
            if isinstance(seedance_prompt_data, dict):
                seedance_prompt = seedance_prompt_data.get('seedance_prompt', '')
                product_images = seedance_prompt_data.get('product_images', [])
            else:
                seedance_prompt = seedance_prompt_data
                product_images = []
            
            ws_scenes.cell(row=i, column=9, value=seedance_prompt)
            
            if not clean_mode:
                if product_images:
                    ws_scenes.row_dimensions[i].height = max(100, len(product_images) * 80)
                    
                    for img_idx, img_path in enumerate(product_images):
                        if os.path.exists(img_path):
                            try:
                                img = OpenpyxlImage(img_path)
                                img.width = 150
                                img.height = 150
                                
                                cell_position = f'J{i}'
                                img.anchor = cell_position
                                ws_scenes.add_image(img)
                                
                                ws_scenes.cell(row=i, column=10, value=os.path.basename(img_path))
                            except Exception as e:
                                print(f"无法插入图片 {img_path}: {e}")
                
                relative_path = os.path.relpath(scene['keyframe_path'], self.output_dir)
                ws_scenes.cell(row=i, column=11, value=f"./{relative_path}")
        
        ws_scenes.column_dimensions['A'].width = 12
        ws_scenes.column_dimensions['B'].width = 15
        ws_scenes.column_dimensions['C'].width = 15
        ws_scenes.column_dimensions['D'].width = 12
        ws_scenes.column_dimensions['E'].width = 15
        ws_scenes.column_dimensions['F'].width = 60
        ws_scenes.column_dimensions['G'].width = 50
        ws_scenes.column_dimensions['H'].width = 100
        ws_scenes.column_dimensions['I'].width = 100
        if not clean_mode:
            ws_scenes.column_dimensions['J'].width = 30
            ws_scenes.column_dimensions['K'].width = 40
        
        for row in ws_scenes.iter_rows(min_row=2, max_row=len(scenes)+1):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        excel_path = os.path.join(self.output_dir, 'video_script.xlsx')
        wb.save(excel_path)
        wb.close()
        
        return excel_path
    
    def extract_shot_type(self, prompt):
        prompt_lower = prompt.lower()
        if 'close-up' in prompt_lower or 'close up' in prompt_lower:
            return '特写'
        elif 'medium shot' in prompt_lower or 'mid shot' in prompt_lower:
            return '中景'
        elif 'wide shot' in prompt_lower or 'long shot' in prompt_lower or 'full shot' in prompt_lower:
            return '全景'
        elif 'extreme close-up' in prompt_lower:
            return '大特写'
        else:
            return '中景'
