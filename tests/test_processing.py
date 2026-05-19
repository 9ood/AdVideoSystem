import os
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch

import app as app_module
from ai_analyzer import AIAnalyzer, AIServiceError
from excel_generator import ExcelGenerator
from openpyxl import load_workbook


class FakeResponse:
    def __init__(self, status_code, payload):
        self.status_code = status_code
        self._payload = payload
        self.ok = 200 <= status_code < 300
        self.text = str(payload)

    def json(self):
        return self._payload


class InvalidJsonResponse:
    def __init__(self, text="not-json"):
        self.status_code = 200
        self.ok = True
        self.text = text

    def json(self):
        raise ValueError("invalid json")


class TimeoutThenSuccess:
    def __init__(self):
        self.calls = 0

    def __call__(self, *args, **kwargs):
        self.calls += 1
        if self.calls == 1:
            raise __import__("requests").Timeout("timed out")
        return FakeResponse(
            200,
            {
                "choices": [{"message": {"content": "ok after timeout retry"}}],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1},
            },
        )


class FakeVideoProcessor:
    def __init__(self, video_path, output_dir):
        self.video_path = video_path
        self.output_dir = output_dir

    def detect_scenes(self):
        return [(0, 1)]

    def extract_keyframes(self, scenes):
        frame_path = os.path.join(self.output_dir, "scene_001_keyframe.jpg")
        with open(frame_path, "wb") as handle:
            handle.write(b"fake")
        return [
            {
                "scene_number": 1,
                "start_time": 0.0,
                "end_time": 1.0,
                "duration": 1.0,
                "keyframe_path": frame_path,
                "first_frame_path": frame_path,
                "last_frame_path": frame_path,
            }
        ]


class FailingAnalyzer:
    def __init__(self, api_key, model, base_url=None):
        self.api_key = api_key
        self.model = model
        self.base_url = base_url

    def analyze_frame(self, image_path):
        raise AIServiceError("分析镜头画面失败：模型不可用")

    def analyze_creative_core_from_keyframes(self, frames_info):
        return {}


class ProcessFlowTests(unittest.TestCase):
    def test_ai_analyzer_falls_back_to_available_model(self):
        forbidden = FakeResponse(
            403,
            {"error": {"message": "This model is not available in your region.", "code": 403}},
        )
        success = FakeResponse(
            200,
            {
                "choices": [{"message": {"content": "ok from fallback"}}],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1},
            },
        )

        with patch.dict(
            os.environ,
            {"AI_FALLBACK_MODELS": "qwen/qwen2.5-vl-72b-instruct"},
            clear=False,
        ):
            analyzer = AIAnalyzer("test-key", "google/gemini-3-flash-preview")
            with patch.object(analyzer, "encode_image", return_value="abc"), patch(
                "ai_analyzer.requests.post",
                side_effect=[forbidden, success],
            ), patch("ai_analyzer._log_token_usage"):
                result = analyzer.analyze_frame("unused.jpg")

        self.assertEqual(result, "ok from fallback")
        self.assertEqual(analyzer.model, "qwen/qwen2.5-vl-72b-instruct")

    def test_process_video_returns_error_when_ai_service_fails(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            upload_dir = os.path.join(temp_dir, "uploads")
            output_dir = os.path.join(temp_dir, "outputs")
            os.makedirs(upload_dir, exist_ok=True)
            os.makedirs(output_dir, exist_ok=True)

            video_path = os.path.join(upload_dir, "task_demo.mp4")
            with open(video_path, "wb") as handle:
                handle.write(b"fake video")

            app_module.app.config["TESTING"] = True
            app_module.app.config["UPLOAD_FOLDER"] = upload_dir
            app_module.app.config["OUTPUT_FOLDER"] = output_dir

            with app_module.app.test_client() as client, patch.object(
                app_module,
                "get_task_video",
                return_value=("demo.mp4", video_path),
            ), patch.object(
                app_module,
                "VideoProcessor",
                FakeVideoProcessor,
            ), patch.object(
                app_module,
                "AIAnalyzer",
                FailingAnalyzer,
            ), patch.dict(
                os.environ,
                {"OPENROUTER_API_KEY": "test-key"},
                clear=False,
            ):
                response = client.post("/process/task-demo")

        self.assertEqual(response.status_code, 502)
        data = response.get_json()
        self.assertIn("模型不可用", data["error"])

    def test_ai_analyzer_retries_when_provider_returns_invalid_json(self):
        bad = InvalidJsonResponse(" \n \n provider glitch")
        success = FakeResponse(
            200,
            {
                "choices": [{"message": {"content": "ok after retry"}}],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1},
            },
        )

        analyzer = AIAnalyzer("test-key", "qwen/qwen3.6-plus")
        with patch.object(analyzer, "encode_image", return_value="abc"), patch(
            "ai_analyzer.requests.post",
            side_effect=[bad, success],
        ), patch("ai_analyzer._log_token_usage"):
            result = analyzer.analyze_frame("unused.jpg")

        self.assertEqual(result, "ok after retry")
        self.assertEqual(analyzer.model, "qwen/qwen3.6-plus")

    def test_qwen_openrouter_requests_disable_reasoning_by_default(self):
        analyzer = AIAnalyzer(
            "test-key",
            "qwen/qwen3.6-plus",
            base_url="https://openrouter.ai/api/v1/chat/completions",
        )

        payload = analyzer._build_request_payload(
            "qwen/qwen3.6-plus",
            [{"role": "user", "content": "hello"}],
        )

        self.assertEqual(payload["reasoning"]["effort"], "none")
        self.assertTrue(payload["reasoning"]["exclude"])

    def test_non_qwen_models_do_not_force_reasoning_config(self):
        analyzer = AIAnalyzer(
            "test-key",
            "google/gemini-3-flash-preview",
            base_url="https://openrouter.ai/api/v1/chat/completions",
        )

        payload = analyzer._build_request_payload(
            "google/gemini-3-flash-preview",
            [{"role": "user", "content": "hello"}],
        )

        self.assertNotIn("reasoning", payload)

    def test_generate_script_messages_do_not_repeat_image_input(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")

        messages = analyzer._build_generate_script_messages("A person speaks to camera.")

        self.assertEqual(messages[0]["content"][0]["type"], "text")
        self.assertEqual(len(messages[0]["content"]), 1)

    def test_ai_analyzer_retries_after_timeout(self):
        retry_post = TimeoutThenSuccess()
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")

        with patch(
            "ai_analyzer.requests.post",
            side_effect=retry_post,
        ), patch("ai_analyzer._log_token_usage"), patch("ai_analyzer.time.sleep"):
            result = analyzer.generate_script("unused.jpg", "A person speaks to camera.")

        self.assertEqual(result, "ok after timeout retry")
        self.assertEqual(retry_post.calls, 2)

    def test_generate_script_normalizes_single_line_output(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")

        with patch.object(
            analyzer,
            "_request_chat_completion",
            return_value="台词：今天这款产品真的很适合刚开始学琴的小朋友。",
        ):
            result = analyzer.generate_script("unused.jpg", "A person speaks to camera.")

        self.assertEqual(result, "今天这款产品真的很适合刚开始学琴的小朋友。")

    def test_generate_script_returns_empty_for_no_dialogue(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")

        with patch.object(
            analyzer,
            "_request_chat_completion",
            return_value="无台词",
        ):
            result = analyzer.generate_script("unused.jpg", "No one is speaking.")

        self.assertEqual(result, "")

    def test_analyze_creative_core_extracts_original_ad_playbook(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")
        captured = {}

        def fake_request(messages, action):
            captured["prompt"] = messages[0]["content"][0]["text"]
            return """1. 一句话创意内核: 孩子像玩游戏一样主动冲去练琴
2. 广告类型: 游戏化奇幻
开头钩子: 放学冲进家门直奔钢琴
故事线: 冲进家门 -> 打开App -> 进入奇幻音乐世界 -> 下载转化
核心冲突/看点: 练琴不再枯燥
App在故事里的角色: 传送门
App界面出现方式: iPad特写和3D UI
证明App有用的方式: 星星奖励和即时反馈
关键镜头套路: 游戏化特效
情绪变化: 兴奋到沉浸
翻拍必须保留的骨架: 主动冲向钢琴和奇幻世界
翻拍可以替换的元素: 人物和房间
最容易被错误套模板的地方: 不能改成妈妈劝孩子练琴
西西魔法钢琴翻拍建议: 保留游戏化玩法并替换为西西魔法钢琴"""

        with patch.object(analyzer, "_request_chat_completion", side_effect=fake_request):
            result = analyzer.analyze_creative_core([
                {
                    "start_time": 0,
                    "end_time": 3,
                    "prompt": "A child runs home and opens a piano learning app.",
                    "script": "",
                }
            ])

        self.assertEqual(result["ad_type"], "游戏化奇幻")
        self.assertIn("主动冲向钢琴", result["must_preserve"])
        self.assertIn("不要默认写成", captured["prompt"])

    def test_analyze_creative_core_from_keyframes_uses_original_images(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")
        captured = {}

        with tempfile.TemporaryDirectory() as temp_dir:
            frame_path = os.path.join(temp_dir, "frame.jpg")
            with open(frame_path, "wb") as handle:
                handle.write(b"fake image")

            def fake_request(messages, action):
                captured["action"] = action
                captured["content"] = messages[0]["content"]
                return """1. 一句话创意内核: 跑步不是跑步，而是给人目标感
2. 广告类型: 运动生活方式
开头钩子: 电子闹钟和NRC界面
故事线: 时间提醒 -> App启动 -> 跑步训练 -> 目的感文案
核心冲突/看点: 自律与目标感
App在故事里的角色: 训练记录工具
App界面出现方式: 手机特写
证明App有用的方式: 训练节奏和打卡
关键镜头套路: 时间跳跃和运动蒙太奇
情绪变化: 困倦到坚定
翻拍必须保留的骨架: 时间、App、训练、目的感字幕
翻拍可以替换的元素: 运动项目和人物
最容易被错误套模板的地方: 不能提前改成钢琴App
西西魔法钢琴翻拍建议: 保留时间打卡和目标感"""

            with patch.object(analyzer, "_request_chat_completion", side_effect=fake_request):
                result = analyzer.analyze_creative_core_from_keyframes([
                    {
                        "start_time": 0,
                        "end_time": 2,
                        "keyframe_path": frame_path,
                    }
                ])

        prompt = captured["content"][0]["text"]
        image_parts = [part for part in captured["content"] if part["type"] == "image_url"]
        self.assertEqual(captured["action"], "根据原始关键帧提取创意内核")
        self.assertEqual(result["ad_type"], "运动生活方式")
        self.assertIn("前13个字段禁止出现“西西魔法钢琴”", prompt)
        self.assertEqual(len(image_parts), 1)

    def test_excel_generator_writes_creative_core_sheet(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generator = ExcelGenerator(temp_dir)
            excel_path = generator.generate(
                {"style": "demo", "main_character": "child", "main_scene": "home"},
                [
                    {
                        "scene_number": 1,
                        "start_time": 0.0,
                        "end_time": 1.0,
                        "duration": 1.0,
                        "prompt": "wide shot",
                        "script": "",
                        "complete_prompt": "",
                        "seedance_prompt": "",
                        "keyframe_path": os.path.join(temp_dir, "missing.jpg"),
                    }
                ],
                creative_core={
                    "one_sentence": "保留原视频玩法",
                    "ad_type": "Vlog挑战",
                    "must_preserve": "倒计时和自拍视频",
                },
                clean_mode=True,
            )

            with open(excel_path, "rb") as handle:
                workbook = load_workbook(BytesIO(handle.read()), read_only=True)
            try:
                self.assertIn("创意内核", workbook.sheetnames)
                self.assertEqual(workbook["创意内核"]["B2"].value, "保留原视频玩法")
            finally:
                workbook.close()
                if hasattr(workbook, "_archive"):
                    workbook._archive.close()

    def test_condensed_script_prompt_locks_creative_playbook(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")
        captured = {}

        def fake_request(messages, action):
            captured["prompt"] = messages[0]["content"][0]["text"]
            return """core story: Piano practice keeps the source ad mechanism.
key scene 1: alarm clock and app trigger
key scene 2: countdown and training montage
key scene 3: purpose slogan ending
creative inheritance: kept time jumps, countdown, montage, and value slogan
complete prompt: A fast montage keeps the original playbook while replacing running with piano practice."""

        creative_core = {
            "opening_hook": "alarm clock",
            "storyline": "time -> app -> training -> slogan",
            "visual_structure": "time jumps, countdown, montage",
            "product_role": "training trigger",
            "proof_method": "action starts after countdown",
            "emotion_curve": "pressure to purpose",
            "must_preserve": "time, app, training, purpose slogan",
            "template_risk": "do not make it a static family practice story",
        }

        with patch.object(analyzer, "_request_chat_completion", side_effect=fake_request):
            result = analyzer.generate_condensed_script(
                [
                    {
                        "start_time": 0,
                        "end_time": 2,
                        "duration": 2,
                        "prompt": "alarm clock and running app",
                        "script": "",
                    }
                ],
                [],
                creative_core=creative_core,
            )

        self.assertIn("CREATIVE LOCK - MUST FOLLOW", captured["prompt"])
        self.assertIn("Do not collapse the remake into one static room", captured["prompt"])
        self.assertIn("关键画面1", captured["prompt"])
        self.assertIn("time jumps, countdown, montage", captured["prompt"])
        self.assertEqual(result["key_scene_2"], "countdown and training montage")
        self.assertEqual(
            result["creative_inheritance_check"],
            "kept time jumps, countdown, montage, and value slogan",
        )

    def test_condensed_seedance_prompt_keeps_creative_lock(self):
        analyzer = AIAnalyzer("test-key", "gemini-3.1-flash-lite-preview")
        captured = {}

        def fake_request(messages, action):
            captured["prompt"] = messages[0]["content"][0]["text"]
            return "seedance prompt"

        with patch.object(analyzer, "_request_chat_completion", side_effect=fake_request):
            result = analyzer.generate_seedance_prompt_for_condensed(
                {
                    "core_story": "demo",
                    "complete_prompt": "demo",
                },
                creative_core={
                    "visual_structure": "fake documentary",
                    "must_preserve": "director loses control of the shoot",
                    "template_risk": "do not turn it into normal family practice",
                },
            )

        self.assertEqual(result["seedance_prompt"], "seedance prompt")
        self.assertIn("CREATIVE LOCK - MUST FOLLOW", captured["prompt"])
        self.assertIn("fake documentary", captured["prompt"])
        self.assertIn("director loses control", captured["prompt"])


if __name__ == "__main__":
    unittest.main()
